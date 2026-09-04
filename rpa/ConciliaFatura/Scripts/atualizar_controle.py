from pathlib import Path
import csv
import openpyxl

BASE = Path(r"C:\Users\jeova\Downloads\Automação de Conciliação de Faturas e Pedidos")
CSV_PATH = BASE / "rpa" / "ConciliaFatura" / "Data" / "Relatorios" / "relatorio_conciliacao.csv"
XLSX_PATH = BASE / "dados" / "resultados" / "controle_conciliacao.xlsx"

with CSV_PATH.open("r", encoding="utf-8-sig", newline="") as f:
    rows = list(csv.DictReader(f))

if not rows:
    raise SystemExit("CSV sem dados para gravar no controle.")

item = rows[-1]
wb = openpyxl.load_workbook(XLSX_PATH)
ws = wb["Controle"]

next_row = ws.max_row + 1
while next_row > 2 and all(ws.cell(next_row - 1, col).value in (None, "") for col in range(1, 13)):
    next_row -= 1

ws.append([
    __import__("datetime").datetime.now().isoformat(timespec="seconds"),
    item["numero_fatura"],
    item["numero_pedido"],
    item["cnpj"],
    item["fornecedor"],
    item["valor_fatura"],
    item["valor_pedido"],
    item["diferenca"],
    item["resultado"],
    item["motivo"],
    __import__("getpass").getuser(),
    "0",
])

wb.save(XLSX_PATH)
